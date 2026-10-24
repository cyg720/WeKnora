"""Read-only checks for the local Docker deployment; never prints .env secrets."""
import datetime
import json
import subprocess
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent

def run(args, data=None):
    result = subprocess.run(args, input=data, text=True, encoding='utf-8',
                            errors='replace', capture_output=True, cwd=ROOT, timeout=90)
    if result.returncode:
        raise RuntimeError(result.stderr.strip()[:1200])
    return result.stdout.strip()

def http(url):
    with urllib.request.urlopen(url, timeout=20) as response:
        return {'status': response.status, 'body': response.read().decode('utf-8')}

checks = []
# Compose prints JSONL on current versions.
raw = run(['docker', 'compose', 'ps', '--format', 'json'])
containers = json.loads(raw) if raw.startswith('[') else [json.loads(line) for line in raw.splitlines() if line]
assert {c['Service'] for c in containers} == {'app', 'frontend', 'postgres', 'redis', 'docreader'}
assert all(c['State'] == 'running' for c in containers)
for c in containers:
    if c['Service'] in {'app', 'postgres', 'docreader'}:
        assert c.get('Health') == 'healthy', c['Service']
    checks.append({'name': c['Service'], 'status': 'pass', 'health': c.get('Health') or 'checked separately'})

for url in ['http://localhost/', 'http://localhost:8080/health', 'http://localhost/api/v1/auth/config']:
    result = http(url)
    assert result['status'] == 200
    if url.endswith('/health'):
        assert json.loads(result['body'])['status'] == 'ok'
    if url.endswith('/auth/config'):
        assert json.loads(result['body'])['success'] is True
    checks.append({'name': url, 'status': 'pass'})

db = run(['docker', 'compose', 'exec', '-T', 'postgres', 'sh', '-c',
          'psql -U "$POSTGRES_USER" -d "$POSTGRES_DB" -At -c "SELECT current_database(); SELECT extname FROM pg_extension; SELECT count(*) FROM information_schema.tables WHERE table_schema=\'public\';"'])
assert 'WeKnora' in db
assert 'vector' in db.splitlines() and 'pg_search' in db.splitlines()
assert int(db.splitlines()[-1]) > 0, 'Database migrations did not create tables'
checks.append({'name': 'PostgreSQL database/extensions/tables', 'status': 'pass', 'result': db})

env = dict(line.split('=', 1) for line in (ROOT/'.env').read_text(encoding='utf-8').splitlines()
           if line and not line.startswith('#') and '=' in line)
reply = run(['docker', 'compose', 'exec', '-T', '-e', 'REDISCLI_AUTH='+env['REDIS_PASSWORD'], 'redis', 'redis-cli', 'ping'])
assert reply == 'PONG'
checks.append({'name': 'Redis authenticated PING', 'status': 'pass'})

parser_test = '''import io,json,grpc
from docx import Document
from docreader.proto import docreader_pb2 as pb,docreader_pb2_grpc as rpc
doc=Document();doc.add_paragraph('WeKnora startup document parsing check.')
buf=io.BytesIO();doc.save(buf)
with grpc.insecure_channel('localhost:50051') as channel:
 stub=rpc.DocReaderStub(channel)
 engines=stub.ListEngines(pb.ListEnginesRequest(),timeout=20)
 result=stub.Read(pb.ReadRequest(file_content=buf.getvalue(),file_name='startup-check.docx',file_type='docx',request_id='startup-check'),timeout=60)
 assert not result.error,result.error
 assert 'WeKnora startup document parsing check' in result.markdown_content,result.markdown_content
 print(json.dumps({'engines':[e.name for e in engines.engines],'docxParse':'pass'}))
'''
parsed = json.loads(run(['docker', 'compose', 'exec', '-T', 'docreader', 'python', '-'], parser_test))
checks.append({'name': 'Docreader gRPC discovery and real DOCX parse', 'status': 'pass', 'result': parsed})

images = []
for name in ['WeKnora-app', 'WeKnora-frontend', 'WeKnora-docreader', 'WeKnora-postgres', 'WeKnora-redis']:
    digest = run(['docker', 'inspect', name, '--format', '{{.Image}}'])
    meta = json.loads(run(['docker', 'image', 'inspect', digest]))[0]
    images.append({'container': name, 'id': digest, 'repoDigests': meta.get('RepoDigests'), 'labels': meta['Config'].get('Labels')})
report = {'status': 'pass', 'generated': datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).isoformat(),
          'checks': checks, 'images': images, 'modelValidation': 'Pending user model endpoint/API credentials; no provider configured or calls made.'}
OUT.mkdir(parents=True, exist_ok=True)
(OUT/'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'status': 'pass', 'checks': len(checks)}, ensure_ascii=False))
